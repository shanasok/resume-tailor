# src/resume_parser.py — .docx text extraction (used by the Base Résumé import flow) and
# the candidate's chat-captured "additional experience" file that base_resume merges into
# the résumé text. The base résumé itself now lives in the per-user DB profile (src/base_resume).

from pathlib import Path                 # clean, OS-independent way to point at files (see PRIMER §14)
from docx import Document                # the python-docx library: reads/writes Word .docx files

# The resumes/ folder at the project root. Path(__file__) is THIS file; .parent.parent
# goes up two folders (src/ → project root).
RESUMES_DIR = Path(__file__).parent.parent / "resumes"

# Candidate-provided experience added via the Chat feature. Appended to the base resume
# so every future generation can draw on it. Plain text, one bullet per line.
ADDITIONS_PATH = RESUMES_DIR / "additional_experience.txt"


def extract_docx_text(source) -> str:
    """All the text inside a .docx — paragraphs first, then table cells — joined with
    newlines. `source` can be a Path OR an open file-like object (python-docx takes both),
    which lets the upload route validate a file's bytes BEFORE saving them anywhere."""
    doc = Document(source)               # open the Word document
    lines = []                           # we'll collect every non-empty line of text here

    # A .docx is made of "paragraphs". Loop over them and keep the ones with actual text.
    for para in doc.paragraphs:
        text = para.text.strip()         # .strip() removes leading/trailing whitespace
        if text:                         # skip blank paragraphs (empty string is "falsy")
            lines.append(text)           # add this line to our list

    # Some résumé content can live inside TABLES (e.g. a skills grid), not paragraphs.
    # Loop table → row → cell to pull that text out too. (Three nested loops.)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    lines.append(text)

    return "\n".join(lines)              # join all the lines into one big string


def append_additions(text: str) -> str:
    """Append the candidate's chat-captured extra experience to some résumé text, under a
    clear header so the AI treats it as part of the base résumé (but still all true). A no-op
    when there are none. Shared by the .docx path and the DB-profile path (src/base_resume)."""
    extras = read_additions()
    if extras:
        text += ("\n\nADDITIONAL EXPERIENCE (candidate-provided, all true — treat as part "
                 "of the base resume):\n" + extras)
    return text


def read_additions() -> str:
    """Return the candidate's chat-added experience (empty string if the file doesn't exist)."""
    if ADDITIONS_PATH.exists():                                  # is there an additions file?
        return ADDITIONS_PATH.read_text(encoding="utf-8").strip()  # read it and trim whitespace
    return ""                                                    # none yet → empty string


def append_addition(bullet: str) -> None:
    """Append one real experience bullet (provided by the candidate via chat) to the
    additions file. '-> None' means it doesn't return anything — it just writes to disk."""
    # Clean up the incoming text: trim spaces, strip any leading "-" or "•" bullet marks, trim again.
    bullet = bullet.strip().lstrip("-• ").strip()
    if not bullet:                       # nothing left after cleaning → do nothing
        return
    # Make sure the resumes/ folder exists (it always should, but this is safe).
    ADDITIONS_PATH.parent.mkdir(parents=True, exist_ok=True)
    # If the file already has content, put a newline before our new line so they don't merge.
    prefix = "\n" if read_additions() else ""
    # Open the file in "append" mode ("a") and write the new bullet. The "with" block makes
    # sure the file is properly closed afterward (see PRIMER §12).
    with ADDITIONS_PATH.open("a", encoding="utf-8") as f:
        f.write(f"{prefix}- {bullet}\n")
