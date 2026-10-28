# tests/test_resume_parser.py — path resolution + docx text extraction.
from io import BytesIO

import pytest
from docx import Document

from src import resume_parser


@pytest.fixture
def resumes_dir(tmp_path, monkeypatch):
    """Point the additions-file location at a temp folder so tests never touch the real
    (personal, gitignored) resumes/ directory."""
    monkeypatch.setattr(resume_parser, "ADDITIONS_PATH", tmp_path / "additional_experience.txt")
    return tmp_path


def test_extract_docx_text_reads_paragraphs_and_tables(tmp_path):
    path = tmp_path / "r.docx"
    doc = Document()
    doc.add_paragraph("Jane Doe")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Python, SQL"
    doc.save(path)
    text = resume_parser.extract_docx_text(path)
    assert "Jane Doe" in text
    assert "Python, SQL" in text


def test_extract_docx_text_accepts_file_like(tmp_path):
    buf = BytesIO()
    doc = Document()
    doc.add_paragraph("From memory")
    doc.save(buf)
    buf.seek(0)
    assert "From memory" in resume_parser.extract_docx_text(buf)


def test_add_experience_falls_back_when_no_base(resumes_dir):
    """resume_editor: with NO base resume, a chat-captured bullet must land in the
    additions file (never crash)."""
    from src import resume_editor
    note = resume_editor.add_experience("Built a CI pipeline", company="Acme")
    assert "additional-experience" in note
    assert "Built a CI pipeline" in resume_parser.read_additions()
